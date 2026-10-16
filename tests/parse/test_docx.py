from pathlib import Path
from typing import Any

import docx
import pytest
from docx.oxml import OxmlElement

from spec_bridge.parse.blocks import ParseError
from spec_bridge.parse.dispatch import parse_document
from tests.parse.goldmatch import dev_documents, recall

DEV_DOCX = dev_documents("docx")


@pytest.mark.parametrize("input_id, path, answer", DEV_DOCX)
def test_every_gold_observation_is_located(input_id: str, path: Path, answer: Any) -> None:
    found, total, missing = recall(parse_document(path, input_id), answer)
    assert found == total, missing


def test_tracked_changes_are_reported(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Capacity ")
    insertion = OxmlElement("w:ins")  # a tracked insertion around the value
    insertion.append(paragraph.add_run("120")._r)
    paragraph._p.append(insertion)
    path = tmp_path / "tracked.docx"
    document.save(str(path))
    with pytest.raises(ParseError, match="tracked changes"):
        parse_document(path, "T-1")


def test_title_headings_and_table_cells_are_marked() -> None:
    input_id, path, _ = next(row for row in DEV_DOCX if row[0] == "SYN-004")
    kinds = {b.kind for b in parse_document(path, input_id).blocks}
    assert {"title", "heading", "paragraph", "table_cell"} <= kinds


def test_paragraph_numbers_follow_document_order() -> None:
    input_id, path, _ = DEV_DOCX[0]
    numbers = [
        b.location.paragraph
        for b in parse_document(path, input_id).blocks
        if b.location.paragraph is not None
    ]
    assert numbers == sorted(numbers) and len(numbers) == len(set(numbers))
