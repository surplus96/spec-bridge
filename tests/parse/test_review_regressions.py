"""CP-S2 review findings (docs/reviews/cp-s2-review-result.md) as regression tests."""

import importlib
import os
import shutil
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import docx
import pytest
from docling.datamodel.base_models import ConversionStatus
from docling_core.types.doc import BoundingBox, CoordOrigin
from openpyxl import Workbook

from spec_bridge.eval.locations import contains_token, same_place
from spec_bridge.parse import pdf_parser, run
from spec_bridge.parse.blocks import Block, ParseError
from spec_bridge.parse.dispatch import parse_document
from spec_bridge.parse.run import run_split
from tests.parse.goldmatch import dev_documents

DEV_XLSX_ID, DEV_XLSX, _ = dev_documents("xlsx")[0]
DEV_PDF_ID, DEV_PDF, DEV_PDF_ANSWER = dev_documents("pdf")[0]


def use_inputs(monkeypatch: pytest.MonkeyPatch, inputs: list[tuple[str, Path]]) -> None:
    monkeypatch.setattr(run, "split_inputs", lambda split: inputs)


# R1: every document-level failure is a ParseError and the batch goes on.
def test_missing_file_is_reported(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="cannot read"):
        parse_document(tmp_path / "missing.pdf", "M-1")


def test_batch_continues_after_a_failed_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    use_inputs(monkeypatch, [("MISSING", tmp_path / "missing.pdf"), (DEV_XLSX_ID, DEV_XLSX)])
    summary = run_split("dev", tmp_path / "out")
    assert [f["input_id"] for f in summary["failures"]] == ["MISSING"]
    assert summary["parsed_ids"] == [DEV_XLSX_ID]
    assert (tmp_path / "out/_summary.json").exists()


# R2: a failed re-parse must not leave the previous success behind.
def test_failed_reparse_removes_stale_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source, out = tmp_path / "CACHE.xlsx", tmp_path / "out"
    shutil.copy(DEV_XLSX, source)
    use_inputs(monkeypatch, [("CACHE", source)])
    assert run_split("dev", out)["parsed_ids"] == ["CACHE"]
    source.write_bytes(b"broken")
    second = run_split("dev", out)
    assert second["parsed_ids"] == [] and not (out / "CACHE.json").exists()
    shutil.copy(DEV_XLSX, source)
    assert run_split("dev", out)["parsed_ids"] == ["CACHE"] and (out / "CACHE.json").exists()


# R3: Docling partial success is not a complete parse.
@pytest.mark.docling
def test_partial_docling_conversion_is_reported(monkeypatch: pytest.MonkeyPatch) -> None:
    real = pdf_parser.converter()

    def partial(path: Path) -> Any:
        result = real.convert(path)
        result.status = ConversionStatus.PARTIAL_SUCCESS
        return result

    monkeypatch.setattr(pdf_parser, "converter", lambda: SimpleNamespace(convert=partial))
    with pytest.raises(ParseError, match="partial_success"):
        parse_document(DEV_PDF, DEV_PDF_ID)


# R4: a formula without a cached value is not silently an empty cell.
def test_uncached_formula_is_rejected(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"], sheet["B1"] = "Value", "=1+1"
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    with pytest.raises(ParseError, match="formula"):
        parse_document(path, "F-1")


# R5: numeric cells keep their type and display format for W3.
def test_number_format_is_kept(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"], sheet["B1"] = "Ratio", 0.125
    sheet["B1"].number_format = "0.0%"
    path = tmp_path / "number.xlsx"
    workbook.save(path)
    cell = next(b for b in parse_document(path, "N-1").blocks if b.location.cell_ref == "B1")
    assert (cell.text, cell.cell_type, cell.number_format) == ("0.125", "n", "0.0%")


# R6: nested tables cannot be located with the current Location; reject them.
def test_nested_table_is_rejected(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Specification")
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "NESTED VALUE"
    path = tmp_path / "nested.docx"
    document.save(str(path))
    with pytest.raises(ParseError, match="nested table"):
        parse_document(path, "N-2")


def test_merged_cell_is_one_block(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "H1", "H2"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "MERGED"
    path = tmp_path / "merged.docx"
    document.save(str(path))
    merged = [b for b in parse_document(path, "M-2").blocks if b.text == "MERGED"]
    assert [(b.location.row, b.location.column) for b in merged] == [(2, 1)]


# Additional: coordinate conversion, offline guard, placement negatives, run metadata.
def test_topleft_box_converts_to_bottom_left() -> None:
    box = BoundingBox(l=10, t=20, r=30, b=40, coord_origin=CoordOrigin.TOPLEFT)
    assert pdf_parser._bottom_left(box, 100.0) == [10, 60, 30, 80]


def test_offline_flag_is_forced(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("HF_HUB_OFFLINE", "0")
    importlib.reload(pdf_parser)
    assert os.environ["HF_HUB_OFFLINE"] == "1"


def test_adjacent_row_is_not_the_same_place() -> None:
    gold = next(o["location"] for o in DEV_PDF_ANSWER["observations"] if o["location"].get("bbox"))
    x0, y0, x1, y1 = gold["bbox"]
    shifted = gold | {"bbox": [x0, y0 - 2 * (y1 - y0), x1, y1 - 2 * (y1 - y0)]}
    block = Block.model_validate(
        {
            "kind": "paragraph",
            "text": "x",
            "location": {"pdf_page": gold["pdf_page"], "bbox": gold["bbox"]},
        }
    )
    assert same_place(gold, block) and not same_place(shifted, block)


def test_number_must_match_as_a_whole_token() -> None:
    assert contains_token("Cylinders 6 —", "6")
    assert not contains_token("Cylinders 16", "6")
    assert not contains_token("17220.6 kW", "6")
    assert contains_token("1400 × 795 × 1005", "1400 × 795 × 1005")


@pytest.mark.docling
def test_run_summary_records_reproducibility_metadata(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    use_inputs(monkeypatch, [(DEV_XLSX_ID, DEV_XLSX)])
    env = run_split("dev", tmp_path)["environment"]
    packages = {"docling", "docling-core", "pypdfium2", "openpyxl", "python-docx"}
    assert packages <= env["versions"].keys()
    assert len(env["uv_lock_sha256"]) == 64 and len(env["models_sha256"]) == 64
    assert env["pdf_options"] == {"do_ocr": False, "do_table_structure": True}
