"""Excel parser (openpyxl): one block per non-empty cell with sheet/cell position."""

import warnings
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell

from spec_bridge.ontology.facts import Location
from spec_bridge.parse.blocks import Block, BlockKind, ParseError


def parse(path: Path, document_id: str) -> list[Block]:
    values = _load(path, data_only=True)  # cached results of formulas
    formulas = _load(path, data_only=False)  # formula text, to detect missing caches
    blocks: list[Block] = []
    for sheet in values.worksheets:
        formula_sheet = formulas[sheet.title]
        section: str | None = None
        title_seen = False
        for row in sheet.iter_rows():
            for cell in row:
                source = formula_sheet[cell.coordinate]
                if source.data_type == "f" and cell.value is None:  # CP-S2 R4
                    raise ParseError(
                        document_id,
                        f"formula without a cached value at {sheet.title}!{cell.coordinate};"
                        " recalculate and save the workbook in Excel",
                    )
            cells = [c for c in row if isinstance(c, Cell) and _text(c)]
            if not cells:
                continue
            table_row = len(cells) > 1
            bold = any(c.font is not None and c.font.b for c in cells)
            for cell in cells:
                text = _text(cell)
                kind: BlockKind
                if table_row:
                    kind = "table_cell"
                elif bold and not title_seen:
                    kind, title_seen = "title", True
                elif bold:
                    kind, section = "heading", text
                else:
                    kind = "paragraph"
                location = Location(sheet=sheet.title, cell_ref=cell.coordinate, section=section)
                blocks.append(
                    Block(
                        kind=kind,
                        text=text,
                        section=section,
                        location=location,
                        cell_type=cell.data_type,  # CP-S2 R5: keep type and display format
                        number_format=cell.number_format,
                    )
                )
    return blocks


def _load(path: Path, *, data_only: bool) -> Workbook:
    with warnings.catch_warnings():
        # Minimal workbooks without a default style trigger a harmless openpyxl notice.
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        return load_workbook(path, data_only=data_only)


def _text(cell: Cell | MergedCell) -> str:
    """Cell text as stored; whitespace-only cells count as empty (not '—')."""
    value = cell.value
    if value is None:
        return ""
    text = str(value)
    return text if text.strip() else ""
