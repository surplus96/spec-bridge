"""Read a BOM table (D-101): an Excel sheet or a Markdown table with the same columns.

    Root model: VX001-001            (Excel: a row "Root model" | VX001-001)
    Manufacturer: <maker name>       (Excel: a row "Manufacturer" | <maker name>)
    | Component | Model | Manufacturer | Quantity |

Nothing is interpreted: the cells are kept as written and the linker identifies them.
A file that does not have this shape is refused with the reason.
"""

from pathlib import Path
from typing import Annotated

import openpyxl
from pydantic import Field

from spec_bridge.ontology.facts import StrictModel

COLUMNS = ("Component", "Model", "Manufacturer", "Quantity")
ROOT_MODEL, ROOT_MAKER = "Root model", "Manufacturer"


class BomError(ValueError):
    """The BOM cannot be read or linked as a whole."""


class BomRow(StrictModel):
    component: str  # the class the BOM says the item is
    model: str
    maker: str
    quantity: Annotated[int, Field(ge=1)]
    line: int  # Markdown line or Excel row, for the reviewer


class BomDocument(StrictModel):
    file_name: str
    root_model: str
    root_maker: str
    rows: list[BomRow]


def _quantity(text: object, line: int) -> int:
    value = str(text).strip()
    if not value.isdigit() or int(value) < 1:
        raise BomError(f"line {line}: quantity {value!r} is not a whole number of at least 1")
    return int(value)


def _row(cells: list[str], line: int) -> BomRow:
    component, model, maker, quantity = cells
    if not (component and model and maker):
        raise BomError(f"line {line}: component, model and manufacturer are required")
    return BomRow(
        component=component,
        model=model,
        maker=maker,
        quantity=_quantity(quantity, line),
        line=line,
    )


def _document(name: str, header: dict[str, str], rows: list[BomRow]) -> BomDocument:
    for key in (ROOT_MODEL, ROOT_MAKER):
        if not header.get(key):
            raise BomError(f"{name}: no {key!r} before the table")
    if not rows:
        raise BomError(f"{name}: no table with the columns {', '.join(COLUMNS)}")
    return BomDocument(
        file_name=name, root_model=header[ROOT_MODEL], root_maker=header[ROOT_MAKER], rows=rows
    )


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def read_markdown(path: Path) -> BomDocument:
    header: dict[str, str] = {}
    rows: list[BomRow] = []
    in_table = False
    for number, line in enumerate(path.read_text().splitlines(), start=1):
        if not line.strip().startswith("|"):
            if in_table:
                break  # the table has ended
            key, sep, value = line.partition(":")
            if sep and key.strip() in (ROOT_MODEL, ROOT_MAKER) and key.strip() not in header:
                header[key.strip()] = value.strip()
            continue
        cells = _cells(line)
        if not in_table:
            if tuple(cells) != COLUMNS:
                raise BomError(f"line {number}: the columns must be {', '.join(COLUMNS)}")
            in_table = True
        elif not all(set(c) <= set("-: ") for c in cells):  # skip the |---| line
            rows.append(_row(cells, number))
    return _document(path.name, header, rows)


def read_excel(path: Path) -> BomDocument:
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    header: dict[str, str] = {}
    rows: list[BomRow] = []
    in_table = False
    for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        cells = ["" if v is None else str(v).strip() for v in values]
        if not any(cells):
            if in_table:
                break
            continue
        if in_table:
            rows.append(_row((cells + [""] * 4)[:4], number))
        elif tuple(cells[:4]) == COLUMNS:
            in_table = True
        elif cells[0] in (ROOT_MODEL, ROOT_MAKER) and len(cells) > 1:
            header.setdefault(cells[0], cells[1])
        elif cells[0] in COLUMNS:
            raise BomError(f"row {number}: the columns must be {', '.join(COLUMNS)}")
    return _document(path.name, header, rows)


def read_bom(path: Path) -> BomDocument:
    if path.suffix == ".xlsx":
        return read_excel(path)
    if path.suffix == ".md":
        return read_markdown(path)
    raise BomError(f"{path.name}: a BOM is an .xlsx or .md table")
